from __future__ import annotations

from pathlib import Path

from knowledge_workbench.errors import MissingDependencyError
from knowledge_workbench.models import ParsedUnit, ParseResult


class PdfParser:
    name = "pypdf"
    version = "1"
    extensions = frozenset({".pdf"})

    def parse(self, path: Path) -> ParseResult:
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise MissingDependencyError(
                "解析 PDF 需要 pypdf；请运行 python -m pip install -e .[documents]"
            ) from exc
        reader = PdfReader(path)
        units = tuple(
            ParsedUnit(text=text, locator={"page": page_number})
            for page_number, page in enumerate(reader.pages, start=1)
            if (text := (page.extract_text() or "").strip())
        )
        return ParseResult(self.name, self.version, units)


class DocxParser:
    name = "python-docx"
    version = "1"
    extensions = frozenset({".docx"})

    def parse(self, path: Path) -> ParseResult:
        try:
            from docx import Document
        except ImportError as exc:
            raise MissingDependencyError(
                "解析 Word 需要 python-docx；请运行 python -m pip install -e .[documents]"
            ) from exc
        document = Document(path)
        headings: list[str] = []
        units: list[ParsedUnit] = []
        for number, paragraph in enumerate(document.paragraphs, start=1):
            text = paragraph.text.strip()
            if not text:
                continue
            style_name = paragraph.style.name if paragraph.style else ""
            if style_name.lower().startswith("heading"):
                try:
                    level = int(style_name.rsplit(" ", 1)[-1])
                except ValueError:
                    level = 1
                headings[:] = headings[: level - 1]
                headings.append(text)
                continue
            units.append(
                ParsedUnit(text, {"paragraph": number, "heading_path": headings.copy()})
            )
        return ParseResult(self.name, self.version, tuple(units))


class XlsxParser:
    name = "openpyxl"
    version = "1"
    extensions = frozenset({".xlsx"})

    def parse(self, path: Path) -> ParseResult:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise MissingDependencyError(
                "解析 Excel 需要 openpyxl；请运行 python -m pip install -e .[documents]"
            ) from exc
        workbook = load_workbook(path, read_only=True, data_only=True)
        units: list[ParsedUnit] = []
        try:
            for sheet in workbook.worksheets:
                for row in sheet.iter_rows():
                    populated = [cell for cell in row if cell.value not in (None, "")]
                    if not populated:
                        continue
                    text = " | ".join(f"{cell.coordinate}={cell.value}" for cell in populated)
                    units.append(
                        ParsedUnit(
                            text,
                            {
                                "sheet": sheet.title,
                                "cell_range": f"{populated[0].coordinate}:{populated[-1].coordinate}",
                            },
                        )
                    )
        finally:
            workbook.close()
        return ParseResult(self.name, self.version, tuple(units))


class PptxParser:
    name = "python-pptx"
    version = "1"
    extensions = frozenset({".pptx"})

    def parse(self, path: Path) -> ParseResult:
        try:
            from pptx import Presentation
        except ImportError as exc:
            raise MissingDependencyError(
                "解析 PowerPoint 需要 python-pptx；请运行 python -m pip install -e .[documents]"
            ) from exc
        presentation = Presentation(path)
        units: list[ParsedUnit] = []
        for slide_number, slide in enumerate(presentation.slides, start=1):
            fragments = [
                shape.text.strip()
                for shape in slide.shapes
                if hasattr(shape, "text") and shape.text.strip()
            ]
            if fragments:
                units.append(
                    ParsedUnit("\n".join(fragments), {"slide": slide_number})
                )
        return ParseResult(self.name, self.version, tuple(units))

