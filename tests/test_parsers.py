import importlib.util
import tempfile
import unittest
from pathlib import Path

from knowledge_workbench.errors import UnsupportedFormatError
from knowledge_workbench.parsers import legacy_word, parse_document
from knowledge_workbench.parsers.legacy_word import ConversionMetadata
from knowledge_workbench.parsers.registry import supported_extensions


HAS_DOCUMENT_DEPS = all(
    importlib.util.find_spec(name) for name in ("pypdf", "docx", "openpyxl", "pptx")
)


@unittest.skipUnless(HAS_DOCUMENT_DEPS, "document parser extras are not installed")
class DocumentParserTests(unittest.TestCase):
    def test_word_conversion_script_disables_macros_alerts_and_personal_information(self):
        script = legacy_word._WORD_CONVERSION_SCRIPT

        self.assertIn("$word.DisplayAlerts = 0", script)
        self.assertIn("$word.AutomationSecurity = 3", script)
        self.assertIn("PSObject.Properties['SendPersonalInformation']", script)
        self.assertIn("if ($null -ne $sendPersonalInformation)", script)
        self.assertIn("$word.Options.SendPersonalInformation = $false", script)

    def test_legacy_doc_requires_explicit_conversion_permission(self):
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary) / "sample.doc"
            path.write_bytes(b"legacy word placeholder")

            with self.assertRaisesRegex(
                UnsupportedFormatError, "allow-legacy-word-conversion"
            ):
                parse_document(path)

            self.assertIn(".doc", supported_extensions())

    def test_legacy_doc_conversion_preserves_conversion_provenance(self):
        from docx import Document

        class FakeWordConverter:
            def convert(self, source: Path, output: Path) -> ConversionMetadata:
                self.source = source
                document = Document()
                document.add_heading("转换测试", level=1)
                document.add_paragraph("旧版文件中的可追溯证据。")
                document.save(output)
                return ConversionMetadata("fake-word", "1.2.3")

        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary) / "sample.doc"
            path.write_bytes(b"legacy word placeholder")
            converter = FakeWordConverter()

            result = parse_document(
                path,
                allow_legacy_word_conversion=True,
                legacy_doc_converter=converter,
            )

            self.assertEqual(converter.source, path)
            self.assertEqual(result.parser_name, "fake-word+python-docx")
            self.assertEqual(result.parser_version, "1.2.3/docx-2")
            self.assertEqual(result.units[0].text, "旧版文件中的可追溯证据。")
            self.assertEqual(result.units[0].locator["source_format"], "doc")
            self.assertEqual(result.units[0].locator["converted_format"], "docx")
            self.assertEqual(result.units[0].locator["conversion_tool"], "fake-word")
            self.assertEqual(result.units[0].locator["conversion_tool_version"], "1.2.3")
            self.assertEqual(len(result.units[0].locator["converted_sha256"]), 64)

    def test_sql_is_read_as_traceable_plain_text(self):
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary) / "01-course.sql"
            path.write_text(
                "-- 课程主表\n"
                "CREATE TABLE study_course (\n"
                "  id BIGINT PRIMARY KEY,\n"
                "  title VARCHAR(200) NOT NULL\n"
                ");\n",
                encoding="utf-8",
            )

            result = parse_document(path)

            self.assertIn(".sql", supported_extensions())
            self.assertEqual(result.parser_name, "plain-text")
            self.assertEqual(result.units[0].locator["line_start"], 1)
            self.assertEqual(result.units[0].locator["line_end"], 5)
            self.assertIn("CREATE TABLE study_course", result.units[0].text)

    def test_pdf_page_locator(self):
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary) / "sample.pdf"
            _write_minimal_pdf(path, "Evidence first knowledge")
            result = parse_document(path)
            self.assertIn("Evidence first knowledge", result.units[0].text)
            self.assertEqual(result.units[0].locator["page"], 1)

    def test_docx_heading_and_paragraph_locator(self):
        from docx import Document

        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary) / "sample.docx"
            document = Document()
            document.add_heading("审核规则", level=1)
            document.add_paragraph("每条证据必须保存原文片段。")
            document.save(path)
            result = parse_document(path)
            self.assertEqual(result.units[0].locator["heading_path"], ["审核规则"])
            self.assertEqual(result.units[0].text, "每条证据必须保存原文片段。")

    def test_docx_table_rows_are_preserved_as_traceable_units(self):
        from docx import Document

        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary) / "table.docx"
            document = Document()
            table = document.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "字段"
            table.cell(0, 1).text = "内容"
            table.cell(1, 0).text = "项目名称"
            table.cell(1, 1).text = "知识工作台"
            document.save(path)

            result = parse_document(path)

            self.assertEqual(len(result.units), 2)
            self.assertEqual(result.units[1].text, "C1=项目名称 | C2=知识工作台")
            self.assertEqual(result.units[1].locator["table"], 1)
            self.assertEqual(result.units[1].locator["row"], 2)
            self.assertEqual(result.units[1].locator["cell_range"], "R2C1:R2C2")

    def test_docx_filters_toc_entries_and_infers_numbered_heading(self):
        from docx import Document

        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary) / "numbered-heading.docx"
            document = Document()
            document.add_paragraph("目录")
            document.add_paragraph("1. 文档简介\t4")
            document.add_paragraph("1. 文档简介")
            document.add_paragraph("本文说明证据提取规则。")
            document.save(path)

            result = parse_document(path)

            self.assertEqual(len(result.units), 1)
            self.assertEqual(result.units[0].text, "本文说明证据提取规则。")
            self.assertEqual(result.units[0].locator["heading_path"], ["1. 文档简介"])

    def test_xlsx_sheet_and_cell_range(self):
        from openpyxl import Workbook

        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary) / "sample.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "知识"
            sheet.append(["规则", "必须引用证据"])
            workbook.save(path)
            workbook.close()
            result = parse_document(path)
            self.assertEqual(result.units[0].locator["sheet"], "知识")
            self.assertEqual(result.units[0].locator["cell_range"], "A1:B1")

    def test_pptx_slide_locator(self):
        from pptx import Presentation

        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary) / "sample.pptx"
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[1])
            slide.shapes.title.text = "知识审核"
            slide.placeholders[1].text = "冲突证据不能自动覆盖正式知识"
            presentation.save(path)
            result = parse_document(path)
            self.assertEqual(result.units[0].locator["slide"], 1)
            self.assertIn("冲突证据", result.units[0].text)


def _write_minimal_pdf(path: Path, text: str) -> None:
    escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    stream = f"BT /F1 12 Tf 72 720 Td ({escaped}) Tj ET".encode("ascii")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(stream)).encode("ascii") + b" >>\nstream\n" + stream + b"\nendstream",
    ]
    output = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for number, body in enumerate(objects, start=1):
        offsets.append(len(output))
        output.extend(f"{number} 0 obj\n".encode("ascii"))
        output.extend(body)
        output.extend(b"\nendobj\n")
    xref_offset = len(output)
    output.extend(f"xref\n0 {len(objects) + 1}\n".encode("ascii"))
    output.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        output.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    output.extend(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_offset}\n%%EOF\n".encode(
            "ascii"
        )
    )
    path.write_bytes(output)
